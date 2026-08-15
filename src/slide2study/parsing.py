from __future__ import annotations

import re
from abc import ABC, abstractmethod
from collections import Counter
from math import ceil
from pathlib import Path

from slide2study.models import Page, ParseReport


class DocumentParser(ABC):
    @abstractmethod
    def parse(self, path: str | Path) -> list[Page]:
        """Parse a document into page-level evidence units."""


class TextParser(DocumentParser):
    """Offline-friendly parser; form-feed or `--- page ---` starts a new page."""

    def parse(self, path: str | Path) -> list[Page]:
        source = Path(path)
        _require_file(source)
        raw = source.read_text(encoding="utf-8")
        parts = re.split(r"\f|^\s*---\s*page\s*---\s*$", raw, flags=re.I | re.M)
        return [
            Page(
                source.stem,
                number,
                text.strip(),
                _explicit_text_heading(text),
                {"parser": "text", "source_name": source.name},
            )
            for number, text in enumerate(parts, 1)
            if text.strip()
        ]


class PDFParser(DocumentParser):
    def parse(self, path: str | Path) -> list[Page]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError(
                "PDF support requires: pip install 'slide2study[documents]'"
            ) from exc
        source = Path(path)
        _require_file(source)
        reader = PdfReader(str(source))
        pages = []
        for number, pdf_page in enumerate(reader.pages, 1):
            text = (pdf_page.extract_text() or "").strip()
            box = pdf_page.mediabox
            pages.append(
                Page(
                    source.stem,
                    number,
                    text,
                    metadata={
                        "parser": "pypdf",
                        "source_name": source.name,
                        "width_points": round(float(box.width), 2),
                        "height_points": round(float(box.height), 2),
                        "rotation": int(pdf_page.get("/Rotate", 0) or 0),
                        "image_count": _pdf_image_count(pdf_page),
                    },
                )
            )
        return prepare_pages_for_retrieval(pages)


class PPTXParser(DocumentParser):
    def parse(self, path: str | Path) -> list[Page]:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError(
                "PPTX support requires: pip install 'slide2study[documents]'"
            ) from exc
        source = Path(path)
        _require_file(source)
        deck = Presentation(str(source))
        pages = []
        for number, slide in enumerate(deck.slides, 1):
            ordered_shapes = sorted(
                slide.shapes,
                key=lambda shape: (getattr(shape, "top", 0), getattr(shape, "left", 0)),
            )
            blocks = []
            block_metadata = []
            for shape in ordered_shapes:
                value, block_type = _pptx_shape_text(shape)
                if not value:
                    continue
                blocks.append(value)
                block_metadata.append(
                    {
                        "type": block_type,
                        "text": value,
                        "left": int(getattr(shape, "left", 0)),
                        "top": int(getattr(shape, "top", 0)),
                        "width": int(getattr(shape, "width", 0)),
                        "height": int(getattr(shape, "height", 0)),
                    }
                )
            title_shape = getattr(slide.shapes, "title", None)
            title = title_shape.text.strip() if title_shape is not None else None
            metadata = {
                "parser": "python-pptx",
                "source_name": source.name,
                "blocks": block_metadata,
                "image_count": sum(1 for shape in slide.shapes if _is_picture(shape)),
            }
            notes = _speaker_notes(slide)
            if notes:
                metadata["speaker_notes"] = notes
            pages.append(Page(source.stem, number, "\n".join(blocks), title, metadata))
        return prepare_pages_for_retrieval(pages)


def get_parser(path: str | Path) -> DocumentParser:
    suffix = Path(path).suffix.lower()
    parsers = {".pdf": PDFParser, ".pptx": PPTXParser, ".txt": TextParser, ".md": TextParser}
    try:
        return parsers[suffix]()
    except KeyError as exc:
        raise ValueError(f"Unsupported document type: {suffix or '<none>'}") from exc


def _first_nonempty_line(text: str) -> str | None:
    return next((line.strip() for line in text.splitlines() if line.strip()), None)


def _explicit_text_heading(text: str) -> str | None:
    first = _first_nonempty_line(text)
    return first if first and first.startswith("#") else None


def build_parse_report(pages: list[Page], low_text_threshold: int = 40) -> ParseReport:
    document_id = pages[0].document_id if pages else "unknown"
    empty_pages = [page.page_number for page in pages if not page.text.strip()]
    low_text_pages = [
        page.page_number for page in pages if 0 < len(page.text.strip()) < low_text_threshold
    ]
    image_pages = [page.page_number for page in pages if page.metadata.get("image_count", 0) > 0]
    quality = {page.page_number: _page_quality(page, low_text_threshold) for page in pages}
    requires_vision = [number for number, value in quality.items() if value["requires_vision"]]
    low_value_pages = [
        number
        for number, value in quality.items()
        if value["role"] in {"boilerplate", "section_divider"}
    ]
    excluded_pages = [
        number for number, value in quality.items() if value["text_retrieval_excluded"]
    ]
    page_roles = Counter(value["role"] for value in quality.values())
    warnings = []
    if not pages:
        warnings.append("No pages were parsed")
    if empty_pages:
        warnings.append(f"{len(empty_pages)} page(s) contain no extractable text")
    if low_text_pages:
        warnings.append(f"{len(low_text_pages)} page(s) contain very little text")
    if requires_vision:
        warnings.append(f"{len(requires_vision)} page(s) should use visual understanding")
    if low_value_pages:
        warnings.append(
            f"{len(low_value_pages)} low-value page(s) are excluded from text retrieval"
        )
    if excluded_pages:
        warnings.append(f"{len(excluded_pages)} page(s) are excluded from text retrieval")
    return ParseReport(
        document_id=document_id,
        page_count=len(pages),
        nonempty_pages=len(pages) - len(empty_pages),
        total_characters=sum(len(page.text) for page in pages),
        empty_pages=empty_pages,
        low_text_pages=low_text_pages,
        image_pages=image_pages,
        requires_vision_pages=requires_vision,
        low_value_pages=low_value_pages,
        text_retrieval_excluded_pages=excluded_pages,
        page_roles=dict(sorted(page_roles.items())),
        removed_boilerplate_lines=sum(
            len(page.metadata.get("removed_boilerplate_lines", [])) for page in pages
        ),
        warnings=warnings,
    )


def prepare_pages_for_retrieval(pages: list[Page]) -> list[Page]:
    repeated_lines = _repeated_boilerplate_lines(pages)
    for page in pages:
        original = page.text
        kept_lines = []
        removed_lines = []
        for line in original.splitlines():
            if _normalize_repeated_line(line) in repeated_lines:
                removed_lines.append(line.strip())
            else:
                kept_lines.append(line)
        if removed_lines:
            page.metadata["raw_text"] = original
            page.metadata["removed_boilerplate_lines"] = removed_lines
            page.text = "\n".join(kept_lines).strip()
        page.metadata.update(_page_quality(page, 40))
        if page.metadata["role"] == "section_divider":
            page.metadata["section_break"] = True
    return pages


def _repeated_boilerplate_lines(pages: list[Page]) -> set[str]:
    if len(pages) < 3:
        return set()
    counts: Counter[str] = Counter()
    originals: dict[str, set[str]] = {}
    for page in pages:
        normalized_on_page = set()
        lines = [line for line in page.text.splitlines() if line.strip()]
        boundary_lines = lines[:2] + lines[-2:]
        for line in boundary_lines:
            normalized = _normalize_repeated_line(line)
            if normalized and len(normalized) <= 80:
                normalized_on_page.add(normalized)
                originals.setdefault(normalized, set()).add(line.strip())
        counts.update(normalized_on_page)
    minimum_pages = max(3, ceil(len(pages) * 0.3))
    return {
        line
        for line, count in counts.items()
        if count >= minimum_pages and (_looks_like_footer(line) or len(originals[line]) > 1)
    }


def _normalize_repeated_line(line: str) -> str:
    normalized = re.sub(r"\d+", "<n>", line.strip().lower())
    return re.sub(r"\s+", " ", normalized)


def _looks_like_footer(line: str) -> bool:
    return bool(
        re.search(r"(?:copyright|university|unimelb|lecture|slide|page|<n>\s*/\s*<n>)", line)
    )


def _page_quality(page: Page, low_text_threshold: int) -> dict[str, object]:
    text = page.text.strip()
    lowered = text.lower()
    image_count = int(page.metadata.get("image_count", 0) or 0)
    compact_length = len(re.sub(r"\s+", "", text))
    semantic_lines = [line.strip() for line in text.splitlines() if line.strip()]
    boilerplate_patterns = (
        "copyright act",
        "do not remove this notice",
        "pursuant to part vb",
        "subject to copyright",
    )
    if any(pattern in lowered for pattern in boilerplate_patterns):
        role = "boilerplate"
    elif (
        image_count > 0
        and len(semantic_lines) == 1
        and compact_length <= 100
        and not re.search(r"[.!?。！？]$", semantic_lines[0])
    ):
        role = "section_divider"
    elif image_count > 0 and compact_length < 20:
        role = "visual_only"
    elif not text:
        role = "empty"
    else:
        role = "content"

    reasons = []
    if image_count and compact_length < 120:
        reasons.append("image_with_sparse_text")
    if image_count >= 10:
        reasons.append("many_image_objects")
    if image_count and len(semantic_lines) >= 10 and compact_length < 500:
        reasons.append("dense_visual_layout")
    if role in {"visual_only", "section_divider"}:
        reasons.append(role)
    requires_vision = bool(reasons) and role != "boilerplate"
    excluded = role in {"boilerplate", "visual_only", "section_divider", "empty"}
    return {
        "role": role,
        "visual_risk_score": round(
            min(
                1.0,
                (0.55 if image_count and compact_length < 120 else 0.0)
                + (0.35 if image_count >= 10 else 0.0)
                + (
                    0.45
                    if image_count and len(semantic_lines) >= 10 and compact_length < 500
                    else 0.0
                )
                + (0.1 if compact_length < low_text_threshold else 0.0),
            ),
            2,
        ),
        "visual_risk_reasons": reasons,
        "requires_vision": requires_vision,
        "text_retrieval_excluded": excluded,
    }


def _require_file(path: Path) -> None:
    if not path.is_file():
        raise FileNotFoundError(f"Document does not exist: {path}")


def _pdf_image_count(pdf_page: object) -> int:
    try:
        return len(pdf_page.images)  # type: ignore[attr-defined]
    except Exception:
        return 0


def _pptx_shape_text(shape: object) -> tuple[str, str]:
    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:  # type: ignore[attr-defined]
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        return "\n".join(rows), "table"
    if getattr(shape, "has_text_frame", False):
        return shape.text.strip(), "text"  # type: ignore[attr-defined]
    return "", "unsupported"


def _is_picture(shape: object) -> bool:
    # MSO_SHAPE_TYPE.PICTURE is 13; avoid importing optional pptx at module import time.
    return int(getattr(shape, "shape_type", -1)) == 13


def _speaker_notes(slide: object) -> str:
    try:
        return slide.notes_slide.notes_text_frame.text.strip()  # type: ignore[attr-defined]
    except (AttributeError, ValueError):
        return ""
