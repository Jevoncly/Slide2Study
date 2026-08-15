import importlib.util
import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from slide2study.chunking import HierarchicalChunker, validate_chunk_hierarchy
from slide2study.cli import main as cli_main
from slide2study.evaluation import evaluate, evaluate_page_retrieval, validate_dataset
from slide2study.identifiers import stable_document_id
from slide2study.interfaces import MultimodalPageEncoder
from slide2study.io import write_jsonl
from slide2study.models import Chunk, Page, RenderedPage
from slide2study.parsing import (
    PDFParser,
    PPTXParser,
    TextParser,
    build_parse_report,
    prepare_pages_for_retrieval,
)
from slide2study.retrieval import BM25Retriever, mixed_tokenize
from slide2study.review import build_review_pack
from slide2study.training import mine_hard_negatives
from slide2study.vision import (
    VisualPageRetriever,
    load_page_embedding_cache,
    load_page_manifest,
    render_document,
    write_page_embedding_cache,
    write_page_manifest,
)

ROOT = Path(__file__).resolve().parents[1]


class FakeMultimodalEncoder(MultimodalPageEncoder):
    def __init__(self, *_args, **_kwargs):
        pass

    def encode_pages(self, image_paths: list[str]) -> list[list[float]]:
        return [[1.0, 0.0], [0.0, 1.0]][: len(image_paths)]

    def encode_queries(self, queries: list[str]) -> list[list[float]]:
        return [[0.0, 1.0] for _ in queries]


class BaselineTests(unittest.TestCase):
    def setUp(self):
        pages = TextParser().parse(ROOT / "examples" / "sample_course.txt")
        self.chunks = HierarchicalChunker(max_chars=200).chunk(pages)

    def test_parses_four_pages_with_citations(self):
        self.assertEqual({chunk.page_start for chunk in self.chunks}, {1, 2, 3, 4})

    def test_builds_three_level_cross_page_hierarchy(self):
        by_level = {
            level: [chunk for chunk in self.chunks if chunk.level == level]
            for level in ("section", "page", "passage")
        }
        self.assertEqual(len(by_level["section"]), 2)
        self.assertEqual(len(by_level["page"]), 4)
        self.assertEqual(len(by_level["passage"]), 4)
        self.assertEqual(
            [(chunk.page_start, chunk.page_end) for chunk in by_level["section"]],
            [(1, 3), (4, 4)],
        )
        by_id = {chunk.chunk_id: chunk for chunk in self.chunks}
        for parent in by_level["section"] + by_level["page"]:
            self.assertTrue(parent.child_ids)
            self.assertTrue(
                all(
                    by_id[child_id].parent_id == parent.chunk_id
                    for child_id in parent.child_ids
                )
            )

    def test_chunk_ids_are_stable(self):
        pages = TextParser().parse(ROOT / "examples" / "sample_course.txt")
        repeated = HierarchicalChunker(max_chars=200).chunk(pages)
        self.assertEqual(
            [chunk.chunk_id for chunk in self.chunks],
            [chunk.chunk_id for chunk in repeated],
        )

    def test_document_id_is_content_based_and_rename_independent(self):
        with tempfile.TemporaryDirectory() as directory:
            first = Path(directory) / "lecture.pdf"
            second = Path(directory) / "renamed.pdf"
            first.write_bytes(b"same document")
            second.write_bytes(b"same document")
            self.assertEqual(stable_document_id(first), stable_document_id(second))
            second.write_bytes(b"changed document")
            self.assertNotEqual(stable_document_id(first), stable_document_id(second))

    def test_hierarchy_validator_detects_missing_parent(self):
        orphan = Chunk("orphan", "doc", 1, 1, "text", parent_id="missing")
        self.assertIn("missing parent", validate_chunk_hierarchy([orphan])[0])

    def test_old_chunk_json_remains_compatible(self):
        chunk = Chunk.from_dict(
            {
                "chunk_id": "legacy",
                "document_id": "doc",
                "page_start": 1,
                "page_end": 1,
                "text": "legacy corpus",
            }
        )
        self.assertEqual(chunk.level, "passage")
        self.assertEqual(chunk.child_ids, [])

    def test_chinese_tokenizer_has_bigrams(self):
        self.assertIn("正则", mixed_tokenize("正则化 L2"))
        self.assertIn("l2", mixed_tokenize("正则化 L2"))
        self.assertIn("λ", mixed_tokenize("L(θ)+λ||θ||²"))

    def test_bm25_retrieves_relevant_page(self):
        results = BM25Retriever(self.chunks).search("正则化如何限制模型复杂度", top_k=2)
        self.assertTrue(results)
        self.assertEqual(results[0].chunk.page_start, 2)

    def test_bm25_level_weights_favor_precise_passages(self):
        chunks = [
            Chunk("section", "doc", 1, 1, "gradient descent", level="section"),
            Chunk("page", "doc", 1, 1, "gradient descent", level="page"),
            Chunk("passage", "doc", 1, 1, "gradient descent", level="passage"),
        ]
        results = BM25Retriever(chunks).search("gradient descent", top_k=3)
        self.assertEqual([result.chunk.level for result in results], ["passage", "page", "section"])

    def test_evaluation_metrics(self):
        examples = [{"query": "验证集选择什么", "relevant_pages": [3]}]
        summary = evaluate(BM25Retriever(self.chunks), examples, top_k=3)
        self.assertEqual(summary.recall_at_k, 1.0)
        self.assertEqual(summary.mrr, 1.0)
        self.assertGreater(summary.precision_at_k, 0.0)
        self.assertEqual(summary.no_result_rate, 0.0)
        self.assertLessEqual(summary.ndcg_at_k, 1.0)
        self.assertIn("unlabeled", summary.by_question_type)

    def test_dataset_validation_reports_types_and_rejects_duplicates(self):
        summary = validate_dataset(
            [
                {
                    "id": "q1",
                    "query": "What is regularization?",
                    "relevant_pages": [2],
                    "question_type": "text",
                }
            ],
            require_question_types=True,
        )
        self.assertEqual(summary.question_types, {"text": 1})
        with self.assertRaisesRegex(ValueError, "duplicate id"):
            validate_dataset(
                [
                    {"id": "q1", "query": "first", "relevant_pages": [1]},
                    {"id": "q1", "query": "second", "relevant_pages": [2]},
                ]
            )

    def test_dataset_validation_checks_corpus_references(self):
        document_id = self.chunks[0].document_id
        summary = validate_dataset(
            [
                {
                    "id": "q1",
                    "query": "What is regularization?",
                    "document_id": document_id,
                    "relevant_pages": [2],
                    "question_type": "text",
                    "split": "test",
                }
            ],
            require_question_types=True,
            require_splits=True,
            require_document_ids=True,
            chunks=self.chunks,
        )
        self.assertEqual(summary.splits, {"test": 1})
        self.assertEqual(summary.documents, {document_id: 1})
        with self.assertRaisesRegex(ValueError, "relevant page"):
            validate_dataset(
                [
                    {
                        "id": "q2",
                        "query": "Invalid page",
                        "document_id": document_id,
                        "relevant_pages": [999],
                        "question_type": "text",
                        "split": "test",
                    }
                ],
                chunks=self.chunks,
            )

    def test_cli_validates_full_corpus_before_level_filtering(self):
        page_chunk = next(chunk for chunk in self.chunks if chunk.level == "page")
        example = {
            "id": "q1",
            "query": page_chunk.text,
            "document_id": page_chunk.document_id,
            "relevant_pages": [page_chunk.page_start],
            "relevant_chunk_ids": [page_chunk.chunk_id],
            "question_type": "text",
            "split": "test",
            "annotation_status": "verified",
        }
        with tempfile.TemporaryDirectory() as directory:
            corpus = Path(directory) / "corpus.jsonl"
            dataset = Path(directory) / "dataset.jsonl"
            write_jsonl((chunk.to_dict() for chunk in self.chunks), corpus)
            write_jsonl([example], dataset)
            output = io.StringIO()
            with redirect_stdout(output):
                result = cli_main(
                    [
                        "evaluate",
                        str(corpus),
                        str(dataset),
                        "--levels",
                        "passage",
                        "--split",
                        "test",
                        "--strict-dataset",
                    ]
                )
        self.assertEqual(result, 0)
        report = json.loads(output.getvalue())
        self.assertEqual(report["experiment"]["split"], "test")
        self.assertEqual(report["metrics"]["queries"], 1)

    def test_evaluation_tracks_no_results_types_and_document_scope(self):
        chunks = [
            Chunk("a", "document-a", 1, 1, "shared keyword"),
            Chunk("b", "document-b", 2, 2, "another topic"),
        ]
        examples = [
            {
                "query": "shared keyword",
                "document_id": "document-b",
                "relevant_pages": [1],
                "question_type": "text",
            },
            {
                "query": "missing phrase",
                "relevant_pages": [9],
                "question_type": "visual_only",
            },
        ]
        summary = evaluate(BM25Retriever(chunks), examples, top_k=1)
        self.assertEqual(summary.recall_at_k, 0.0)
        self.assertEqual(summary.no_result_rate, 0.5)
        self.assertEqual(summary.by_question_type["visual_only"]["no_result_rate"], 1.0)
        self.assertGreaterEqual(summary.p95_latency_ms, 0.0)

    def test_hard_negative_mining_excludes_positive(self):
        retriever = BM25Retriever(self.chunks)
        examples = [{"query": "模型评估 数据", "relevant_pages": [3]}]
        triplets = mine_hard_negatives(retriever, examples, self.chunks, top_k=4)
        positive_ids = {chunk.chunk_id for chunk in self.chunks if chunk.page_start == 3}
        self.assertTrue(triplets)
        self.assertTrue(all(item.negative_chunk_id not in positive_ids for item in triplets))

    def test_parse_report_flags_pages_that_need_vision(self):
        pages = [
            Page("deck", 1, "A complete text page " * 5),
            Page("deck", 2, "", metadata={"image_count": 2}),
            Page("deck", 3, "short", metadata={"image_count": 1}),
        ]
        report = build_parse_report(pages, low_text_threshold=20)
        self.assertEqual(report.empty_pages, [2])
        self.assertEqual(report.low_text_pages, [3])
        self.assertEqual(report.requires_vision_pages, [2, 3])
        self.assertAlmostEqual(report.to_dict()["text_coverage"], 2 / 3, places=6)

    def test_repeated_footer_is_removed_without_losing_raw_evidence(self):
        pages = [
            Page("deck", 1, "Introduction\nCourse 2026 1", metadata={"image_count": 1}),
            Page(
                "deck",
                2,
                "Copyright Act notice\nDo not remove this notice\nCourse 2026 2",
                metadata={"image_count": 1},
            ),
            Page(
                "deck",
                3,
                "IPv6 header\nRequired fields and sizes\nCourse 2026 3",
                metadata={"image_count": 1},
            ),
            Page(
                "deck",
                4,
                "Routing algorithms\nDistance vector details\nCourse 2026 4",
                metadata={"image_count": 1},
            ),
        ]
        prepared = prepare_pages_for_retrieval(pages)
        self.assertEqual(prepared[0].text, "Introduction")
        self.assertIn("Course 2026 1", prepared[0].metadata["raw_text"])
        self.assertEqual(prepared[0].metadata["role"], "section_divider")
        self.assertTrue(prepared[0].metadata["section_break"])
        self.assertEqual(prepared[1].metadata["role"], "boilerplate")
        self.assertEqual(prepared[2].metadata["role"], "content")

    def test_low_value_pages_are_preserved_but_not_retrieved(self):
        pages = prepare_pages_for_retrieval(
            [
                Page("deck", 1, "Module 1", metadata={"image_count": 1}),
                Page(
                    "deck",
                    2,
                    "Copyright Act notice\nDo not remove this notice",
                    metadata={"image_count": 1},
                ),
                Page(
                    "deck",
                    3,
                    "IPv6 header\nRequired fields in the header are shown below.",
                    metadata={"image_count": 1},
                ),
            ]
        )
        chunks = HierarchicalChunker().chunk(pages)
        page_chunks = {chunk.page_start: chunk for chunk in chunks if chunk.level == "page"}
        self.assertEqual(set(page_chunks), {1, 2, 3})
        self.assertEqual(page_chunks[1].child_ids, [])
        self.assertEqual(page_chunks[2].child_ids, [])
        self.assertFalse(BM25Retriever(chunks).search("Copyright Act notice"))
        result = BM25Retriever(chunks).search("IPv6 required fields")[0]
        self.assertEqual(result.chunk.page_start, 3)

    def test_dense_visual_layout_is_sent_to_vision_pipeline(self):
        table_lines = [f"row {index} value {index * 2}" for index in range(12)]
        page = Page("deck", 7, "\n".join(table_lines), metadata={"image_count": 1})
        report = build_parse_report([page])
        self.assertEqual(report.requires_vision_pages, [7])

    def test_visual_page_retriever_ranks_cross_modal_embeddings(self):
        pages = [
            RenderedPage("deck", 1, "page-1.png", "deck.pdf", 100, 80, "a"),
            RenderedPage("deck", 2, "page-2.png", "deck.pdf", 100, 80, "b"),
        ]
        results = VisualPageRetriever(pages, FakeMultimodalEncoder()).search("diagram", 2)
        self.assertEqual([result.page.page_number for result in results], [2, 1])
        self.assertEqual(results[0].score, 1.0)

    def test_visual_evaluation_uses_document_scoped_page_labels(self):
        pages = [
            RenderedPage("deck-a", 1, "page-1.png", "a.pdf", 100, 80, "a"),
            RenderedPage("deck-b", 2, "page-2.png", "b.pdf", 100, 80, "b"),
        ]
        retriever = VisualPageRetriever(pages, FakeMultimodalEncoder())
        summary = evaluate_page_retrieval(
            retriever,
            [
                {
                    "query": "diagram",
                    "document_id": "deck-b",
                    "relevant_pages": [2],
                    "question_type": "visual_only",
                }
            ],
            top_k=2,
        )
        self.assertEqual(summary.recall_at_k, 1.0)
        self.assertEqual(summary.mrr, 1.0)

    def test_page_embedding_cache_checks_model_and_image_hash(self):
        page = RenderedPage("deck", 1, "page.png", "deck.pdf", 100, 80, "abc")
        with tempfile.TemporaryDirectory() as directory:
            cache = Path(directory) / "embeddings.json"
            write_page_embedding_cache([page], [[3.0, 4.0]], cache, "fake-clip")
            embeddings, model = load_page_embedding_cache([page], cache, "fake-clip")
            self.assertEqual(model, "fake-clip")
            self.assertAlmostEqual(embeddings[0][0], 0.6)
            changed = RenderedPage("deck", 1, "page.png", "deck.pdf", 100, 80, "changed")
            with self.assertRaisesRegex(ValueError, "image mismatch"):
                load_page_embedding_cache([changed], cache)
            with self.assertRaisesRegex(ValueError, "expected 'other-model'"):
                load_page_embedding_cache([page], cache, "other-model")

    def test_visual_index_and_evaluate_cli_reuse_page_cache(self):
        page_chunk = next(chunk for chunk in self.chunks if chunk.level == "page")
        example = {
            "id": "visual-q1",
            "query": "diagram",
            "document_id": page_chunk.document_id,
            "relevant_pages": [page_chunk.page_start],
            "relevant_chunk_ids": [page_chunk.chunk_id],
            "question_type": "visual_only",
            "split": "test",
            "annotation_status": "verified",
        }
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            manifest = root / "pages.jsonl"
            corpus = root / "corpus.jsonl"
            dataset = root / "dataset.jsonl"
            cache = root / "embeddings.json"
            report_path = root / "report.json"
            page = RenderedPage(
                page_chunk.document_id,
                page_chunk.page_start,
                str(root / "page.png"),
                "lecture.pdf",
                100,
                80,
                "image-hash",
            )
            write_page_manifest([page], manifest)
            write_jsonl((chunk.to_dict() for chunk in self.chunks), corpus)
            write_jsonl([example], dataset)
            with patch(
                "slide2study.cli.SentenceTransformersCLIPEncoder",
                FakeMultimodalEncoder,
            ), redirect_stdout(io.StringIO()):
                self.assertEqual(
                    cli_main(
                        [
                            "visual-index",
                            "--manifests",
                            str(manifest),
                            "--output",
                            str(cache),
                            "--model",
                            "fake-clip",
                        ]
                    ),
                    0,
                )
                self.assertEqual(
                    cli_main(
                        [
                            "visual-evaluate",
                            str(corpus),
                            str(dataset),
                            "--manifests",
                            str(manifest),
                            "--cache",
                            str(cache),
                            "--split",
                            "test",
                            "--top-k",
                            "1",
                            "--output",
                            str(report_path),
                        ]
                    ),
                    0,
                )
            report = json.loads(report_path.read_text(encoding="utf-8"))
        self.assertEqual(report["experiment"]["model"], "fake-clip")
        self.assertEqual(report["metrics"]["recall_at_k"], 1.0)

    def test_page_manifest_round_trip(self):
        page = RenderedPage(
            "deck", 3, "page-0003.png", "deck.pdf", 1280, 720, "abc", True, "content", 0.8
        )
        with tempfile.TemporaryDirectory() as directory:
            manifest = Path(directory) / "pages.jsonl"
            write_page_manifest([page], manifest)
            loaded = load_page_manifest(manifest)
        self.assertEqual(loaded, [page])

    def test_review_pack_copies_evidence_and_builds_export_ui(self):
        page_chunk = next(chunk for chunk in self.chunks if chunk.level == "page")
        example = {
            "id": "review-q1",
            "query": "What evidence is on this page?",
            "document_id": page_chunk.document_id,
            "relevant_pages": [page_chunk.page_start],
            "relevant_chunk_ids": [page_chunk.chunk_id],
            "question_type": "text",
            "split": "test",
            "annotation_status": "candidate",
            "answer_hint": "The page text",
        }
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image = root / "source.png"
            image.write_bytes(b"fake image")
            rendered = RenderedPage(
                page_chunk.document_id,
                page_chunk.page_start,
                str(image),
                "lecture.pdf",
                100,
                80,
                "abc",
            )
            output = root / "review" / "index.html"
            summary = build_review_pack([example], self.chunks, output, [rendered])
            html = output.read_text(encoding="utf-8")
            copied = list((output.parent / "index_assets").glob("*.png"))
        self.assertEqual(summary["examples"], 1)
        self.assertEqual(summary["copied_images"], 1)
        self.assertEqual(len(copied), 1)
        self.assertIn("Slide2Study QA 人工复核", html)
        self.assertIn("导出复核 JSONL", html)

    @unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow is not installed")
    def test_pdf_renderer_writes_stable_page_mapping(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "lecture.pdf"
            source.write_bytes(b"placeholder")
            executable = root / "pdftoppm.exe"
            executable.write_bytes(b"placeholder")

            def fake_run(command: list[str]) -> None:
                prefix = Path(command[-1])
                Image.new("RGB", (320, 180), "white").save(f"{prefix}-1.png")
                Image.new("RGB", (320, 180), "black").save(f"{prefix}-2.png")

            with patch("slide2study.vision._run", side_effect=fake_run):
                pages = render_document(
                    source,
                    root / "rendered",
                    pdftoppm_executable=executable,
                    page_metadata={2: {"requires_vision": True, "visual_risk_score": 0.8}},
                )
            self.assertEqual([page.page_number for page in pages], [1, 2])
            self.assertEqual(Path(pages[0].image_path).name, "page-0001.png")
            self.assertEqual((pages[0].width, pages[0].height), (320, 180))
            self.assertTrue(pages[1].requires_vision)
            self.assertEqual(len(pages[0].sha256), 64)

    @unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow is not installed")
    def test_pptx_renderer_converts_before_page_rendering(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "lecture.pptx"
            source.write_bytes(b"placeholder")
            pdftoppm = root / "pdftoppm.exe"
            soffice = root / "soffice.exe"
            pdftoppm.write_bytes(b"placeholder")
            soffice.write_bytes(b"placeholder")

            def fake_run(command: list[str]) -> None:
                if "--convert-to" in command:
                    output_dir = Path(command[command.index("--outdir") + 1])
                    (output_dir / "lecture.pdf").write_bytes(b"converted")
                else:
                    Image.new("RGB", (320, 180), "white").save(f"{command[-1]}-1.png")

            with patch("slide2study.vision._run", side_effect=fake_run):
                pages = render_document(
                    source,
                    root / "rendered",
                    pdftoppm_executable=pdftoppm,
                    soffice_executable=soffice,
                )
            self.assertEqual(len(pages), 1)
            self.assertEqual(pages[0].source_path, str(source.resolve()))
            self.assertEqual(pages[0].document_id, stable_document_id(source))

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx is not installed")
    def test_pptx_parser_extracts_title_and_table(self):
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "lecture.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = "Optimization"
            table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(5), Inches(2)).table
            table.cell(0, 0).text = "Method"
            table.cell(0, 1).text = "Rate"
            table.cell(1, 0).text = "SGD"
            table.cell(1, 1).text = "O(1/t)"
            deck.save(path)
            page = PPTXParser().parse(path)[0]
        self.assertEqual(page.title, "Optimization")
        self.assertIn("Method | Rate", page.text)
        self.assertIn("SGD | O(1/t)", page.text)
        self.assertTrue(any(block["type"] == "table" for block in page.metadata["blocks"]))

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "pypdf is not installed")
    def test_pdf_parser_records_page_geometry(self):
        from pypdf import PdfWriter

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "blank.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            with path.open("wb") as stream:
                writer.write(stream)
            page = PDFParser().parse(path)[0]
        self.assertEqual(page.metadata["width_points"], 612.0)
        self.assertEqual(page.metadata["height_points"], 792.0)


if __name__ == "__main__":
    unittest.main()
